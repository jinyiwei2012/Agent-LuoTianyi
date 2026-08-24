# -*- coding: utf-8 -*-
"""精简版 READ_IMAGEVIDEO 权限使用场景说明（正文 ≤200 字）
与 gen_acl_doc.py 同目录，生成可直接提交审核的精简说明文档。
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "READ_IMAGEVIDEO_权限使用场景说明_精简版.docx")
ACCENT = RGBColor(0x16, 0x74, 0xA3)
DARK = RGBColor(0x24, 0x34, 0x47)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)

def setf(run, size=12, bold=False, color=DARK, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color

def para(text, size=12, bold=False, center=False, name="宋体"):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); setf(r, size, bold, DARK, name)
    return p

para("权限使用场景说明（精简版）", 16, True, True, "微软雅黑")
para("申请权限：ohos.permission.READ_IMAGEVIDEO", 12, True, True, "微软雅黑")
para("应用：洛天依（HarmonyOS NEXT，v1.0.0）    日期：2026-08-24", 10, False, True)
doc.add_paragraph()

body = ("用户在聊天界面点击「图片」按钮，调用系统相册选择器，从相册选择一张照片。"
        "所选照片仅经用户确认后，以 Base64 通过 HTTPS 加密上传至服务端，"
        "由多模态大模型识别图片内容，洛天依据此回复（如识别美食、风景、宠物）。"
        "本权限仅用于该功能的图片读取：全部由用户主动触发，"
        "不做后台扫描，不读取相册目录，不申请写入/删除权限。"
        "用户可随时在系统权限管理中撤销。")
para(body, 12)
doc.add_paragraph()
para("附：《功能 demo 录屏》记录完整使用流程（首页→聊天页→相册选择器）", 10)

doc.save(OUT)
print("字数:", len(body))
print("已生成:", OUT, os.path.getsize(OUT), "bytes")